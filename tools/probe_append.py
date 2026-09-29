"""Where does openpyxl's `append` actually land?

The rewrite of `write_excel` assumed `ws.append` writes to the next empty
row, starting at row 1 on a fresh sheet. Two tests failed with a blank row
above the header, so that assumption is wrong somewhere. This measures it
instead of reasoning about openpyxl's internals.
"""
import openpyxl

wb = openpyxl.Workbook()
ws = wb.active
print("fresh sheet        : _current_row =", ws._current_row,
      "max_row =", ws.max_row)
ws.append(["a", "b"])
print("after append(hdr)  : _current_row =", ws._current_row,
      "max_row =", ws.max_row)
ws.append([1, 2])
print("after append(data) : max_row =", ws.max_row)
for r in ws.iter_rows(values_only=True):
    print("   row:", r)

print()
wb2 = openpyxl.Workbook()
w2 = wb2.active
w2.cell(row=1, column=1, value="a")
print("after cell(1,1)    : _current_row =", w2._current_row)
w2.append([1, 2])
print("then append        : max_row =", w2.max_row)
for r in w2.iter_rows(values_only=True):
    print("   row:", r)
