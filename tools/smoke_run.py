r"""End-to-end smoke test: generate real files, plan, run, check the output.

Run with:  .venv\Scripts\python tools\smoke_run.py

This is the check that matters most, because it is the only one that starts
from a file on disk and ends at a file on disk, with no mocks anywhere in
between. If this passes, ``aar run`` genuinely runs a pipeline.
"""

from __future__ import annotations

import os
import sys
import tempfile

import pyarrow as pa
import pyarrow.parquet as pq

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from aar.cli import main  # noqa: E402


REGIONS = ["NA", "EU", "APAC"]


def make_data(path: str) -> None:
    """A deterministic orders table, 40 rows, three regions."""
    rows = {
        "id": list(range(1, 41)),
        "region": [REGIONS[i % 3] for i in range(40)],
        "amount": [float(50 * (i % 17) + 25) for i in range(40)],
        "quantity": [(i % 5) + 1 for i in range(40)],
    }
    pq.write_table(pa.table({
        "id": pa.array(rows["id"], type=pa.int64()),
        "region": pa.array(rows["region"], type=pa.string()),
        "amount": pa.array(rows["amount"], type=pa.float64()),
        "quantity": pa.array(rows["quantity"], type=pa.int64()),
    }), path)


PIPELINE = r'''
from aar.sdk import (col, filter_, group_by, lit, parquet, sort, sum_,
                     write_excel, write_parquet)

SRC = r"{src}"
OUT_PARQUET = r"{out_parquet}"
OUT_EXCEL = r"{out_excel}"


def build():
    orders = parquet(SRC, estimated_bytes=2_000_000)
    big = filter_(orders, col("amount") > lit(100.0))
    by_region = group_by(big, "region", aggs={{
        "total": sum_("amount"),
        "orders": sum_("quantity"),
    }})
    ranked = sort(by_region, "total desc")
    # Two writes, one root: the Excel write consumes the Parquet write's
    # result, so the DAG stays a single tree.
    stored = write_parquet(ranked, OUT_PARQUET)
    return write_excel(stored, OUT_EXCEL, sheet="Summary")
'''


def main_run() -> int:
    tmp = tempfile.mkdtemp(prefix="aar_smoke_")
    src = os.path.join(tmp, "orders.parquet")
    out_pq = os.path.join(tmp, "summary.parquet")
    out_xl = os.path.join(tmp, "summary.xlsx")
    make_data(src)

    pipeline_path = os.path.join(tmp, "pipeline.py")
    with open(pipeline_path, "w", encoding="utf-8") as fh:
        fh.write(PIPELINE.format(src=src.replace("\\", "\\\\"),
                                 out_parquet=out_pq.replace("\\", "\\\\"),
                                 out_excel=out_xl.replace("\\", "\\\\")))

    print("=" * 72)
    print("aar explain")
    print("=" * 72)
    rc_explain = main(["explain", pipeline_path])

    print()
    print("=" * 72)
    print("aar run")
    print("=" * 72)
    rc_run = main(["run", pipeline_path])

    print()
    print("=" * 72)
    print("verification")
    print("=" * 72)

    failures = []
    if rc_explain != 0:
        failures.append(f"explain returned {rc_explain}")
    if rc_run != 0:
        failures.append(f"run returned {rc_run}")
    if not os.path.exists(out_pq):
        failures.append("summary.parquet was not written")
    if not os.path.exists(out_xl):
        failures.append("summary.xlsx was not written")

    if os.path.exists(out_pq):
        got = pq.read_table(out_pq)
        print(f"summary.parquet: {got.num_rows} rows, "
              f"columns {got.column_names}")
        totals = {r["region"]: r["total"] for r in got.to_pylist()}
        print(f"totals: {totals}")

        expected = {"NA": 0.0, "EU": 0.0, "APAC": 0.0}
        for i in range(40):
            amount = float(50 * (i % 17) + 25)
            if amount > 100.0:
                expected[REGIONS[i % 3]] += amount
        for region, want in expected.items():
            if abs(totals.get(region, 0.0) - want) > 1e-6:
                failures.append(
                    f"{region}: expected {want}, got {totals.get(region)}")
        values = [t for t in totals.values()]
        if values != sorted(values, reverse=True):
            failures.append("rows are not sorted by total descending")

    if os.path.exists(out_xl):
        import openpyxl
        wb = openpyxl.load_workbook(out_xl, read_only=True)
        print(f"summary.xlsx: sheets {wb.sheetnames}")
        rows = list(wb["Summary"].iter_rows(values_only=True))
        print(f"  header: {rows[0]}")
        print(f"  {len(rows) - 1} data rows")
        wb.close()
        if rows[0] != ("region", "total", "orders"):
            failures.append(f"xlsx header is {rows[0]}, expected "
                            f"('region', 'total', 'orders')")

    print()
    if failures:
        for f in failures:
            print(f"FAIL: {f}")
        return 1
    print("OK: pipeline planned, executed, wrote both files, and the "
          "numbers are correct.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main_run())
